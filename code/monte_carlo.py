#!/usr/bin/env python3
"""
monte_carlo.py -- surrogate Monte Carlo uncertainty analysis of 2060 Scenario B net build-out CAPEX
(Section 4.6, Table 8 and Figure 11 of the paper).

Method (additive response surface, as specified on the workbook's Sensitivity sheet, rows 118-131):
  * six inputs are varied jointly: saturation ceiling gamma (+/-25%, with alpha/beta refit), charger CAPEX per port
    (+/-30%), grid-upgrade cost per kW (+/-50%), EV-share midpoint t50 (+/-5 years), public DC ratio (+/-50%) and
    home-charge share (+/-50%);
  * each input's effect on CAPEX is taken from the one-at-a-time perturbation table (Sensitivity!C120:D125, the low and
    high CAPEX deltas) and sampled uniformly between those two endpoints -- this is linear interpolation between the
    one-at-a-time endpoints, i.e. a surrogate response surface, not a full nonlinear rerun of the workbook;
  * draws are independent; 5,000 trials; numpy default_rng(42);
  * output = baseline CAPEX (Sensitivity!B77) + sum of the six sampled effects; the 90% interval is the 5th-95th
    percentile of the 5,000 outputs.

Usage:  python3 code/monte_carlo.py [workbook.xlsx]      (workbook must be recalculated: open/save in Excel or run recalc)
Writes: outputs/monte_carlo_draws.csv (all 5,000 outputs) and outputs/monte_carlo_summary.csv
"""
import csv, sys
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
WB = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'workbook' / 'Nigeria_EV_Charging_Master_R2.xlsx'
N_TRIALS, SEED = 5000, 42

def load_spec(path=WB):
    """Read baseline CAPEX and the six (low, high) CAPEX effects from the Sensitivity sheet."""
    s = load_workbook(path, data_only=True)['Sensitivity']
    base = s['B77'].value
    if base is None: raise SystemExit('Workbook has no cached values: open and save it in Excel (or run LibreOffice recalc) first.')
    inputs = [(s.cell(row=r, column=1).value, s.cell(row=r, column=3).value, s.cell(row=r, column=4).value) for r in range(120, 126)]
    return base, inputs

def run(base, inputs, n=N_TRIALS, seed=SEED):
    rng = np.random.default_rng(seed); total = np.full(n, float(base)); effects = {}
    for name, lo, hi in inputs:
        e = rng.uniform(lo, hi, n); effects[name] = e; total += e
    return total, effects

if __name__ == '__main__':
    base, inputs = load_spec()
    total, effects = run(base, inputs)
    p5, p50, p95 = np.percentile(total, [5, 50, 95])
    out = ROOT / 'outputs'; out.mkdir(exist_ok=True)
    with open(out / 'monte_carlo_draws.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['trial', 'total_capex_usd_bn'] + [n for n, _, _ in inputs])
        for i in range(len(total)): w.writerow([i + 1, round(total[i], 4)] + [round(effects[n][i], 4) for n, _, _ in inputs])
    with open(out / 'monte_carlo_summary.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['statistic', 'usd_bn'])
        for k, v in [('baseline', base), ('mean', total.mean()), ('median', p50), ('p5', p5), ('p95', p95), ('min', total.min()), ('max', total.max()), ('trials', len(total)), ('seed', SEED)]: w.writerow([k, round(float(v), 3)])
    print(f'Baseline 2060 Scenario B CAPEX: USD {base:.1f} bn')
    for name, lo, hi in inputs: print(f'  {name:<32} effect range [{lo:+7.1f}, {hi:+7.1f}] USD bn')
    print(f'{len(total)} trials (seed {SEED}): median {p50:.1f}, mean {total.mean():.1f}, 90% interval {p5:.1f} to {p95:.1f} USD bn')
