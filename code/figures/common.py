"""common.py -- shared helpers for the figure scripts: locate the recalculated workbook, read series by year, and save.
Row/column positions refer to the authoritative workbook layout (see build_workbook.py header)."""
import os, sys
from pathlib import Path
import numpy as np
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[2]
WB = Path(os.environ.get('MASTER_XLSX', ROOT / 'workbook' / 'Nigeria_EV_Charging_Master_R2.xlsx'))
OUT = Path(os.environ.get('FIG_OUT', ROOT / 'figures')); OUT.mkdir(exist_ok=True)
plt.rcParams.update({'font.size': 11, 'axes.spines.top': False, 'axes.spines.right': False, 'figure.dpi': 300})
COL = {'A': '#d62728', 'B': '#1f77b4', 'C': '#2ca02c'}                      # scenario colours
CLS = ['Motorcycle', 'Private car', 'LCV', 'HGV', 'Light bus', 'Coach']; CCOL = ['#9467bd', '#1f77b4', '#ff7f0e', '#d62728', '#2ca02c', '#8c564b']
YEARS = list(range(2010, 2061)); RY = np.arange(2025, 2061)
wb = load_workbook(WB, data_only=True)
if wb['Summary 2060']['B4'].value is None:
    sys.exit(f'{WB} has no cached values: open and save it in Excel, or run LibreOffice headless recalculation, then rerun.')
fp, asm, sens, lp, dc = wb['Fleet projections'], wb['Assumptions'], wb['Sensitivity'], wb['Load profile'], wb['Delivery check']
R = {s: wb[f'Results Scenario {s}'] for s in 'ABC'}
FP0, RES0, R60 = 18, 4, 39; CLASS_HDR = {'A': 75, 'B': 130, 'C': 185}
SENS_LOW_ROWS = [100, 102, 104, 106, 108, 110, 112, 114, 132]; MC_ROWS = range(120, 126); SENS_BASE = 'B77'
def fpr(y): return FP0 + (y - 2010)
def fleet(s): return np.array([fp.cell(row=fpr(y), column={'A': 3, 'B': 4, 'C': 5}[s]).value for y in YEARS], float)
def cls_stock(s):
    h = CLASS_HDR[s]; return np.array([[fp.cell(row=h + 1 + (y - 2010), column=3 + i).value for y in YEARS] for i in range(6)], float)
def rcol(s, c): return np.array([R[s].cell(row=RES0 + (y - 2025), column=c).value for y in RY], float)
def r60(s, c): return R[s].cell(row=R60, column=c).value
def save(fig, name):
    fig.tight_layout(); fig.savefig(OUT / f'{name}.png', bbox_inches='tight'); fig.savefig(OUT / f'{name}.svg', bbox_inches='tight'); plt.close(fig); print('wrote', OUT / f'{name}.png')
